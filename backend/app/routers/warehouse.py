"""Warehouse management API."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Header, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppException, ErrorCode, ValidationError
from app.core.permissions import Permission, has_permission, require_permission
from app.core.rate_limit import get_client_ip
from app.database import get_db
from app.models.user import User
from app.schemas.warehouse import (
    CountActionRequest,
    CountCreate,
    CountLineResponse,
    CountLinesUpdate,
    CountListResponse,
    CountResponse,
    DocumentListResponse,
    DocumentResponse,
    InboundCreate,
    LedgerListResponse,
    LocationCreate,
    LocationResponse,
    LocationUpdate,
    MaterialCreate,
    MaterialResponse,
    MaterialUpdate,
    OpeningImportResult,
    OutboundCreate,
    PeriodActionRequest,
    PeriodResponse,
    ProjectCreate,
    ProjectResponse,
    ProjectUpdate,
    QrPayloadResponse,
    RebuildRequest,
    RebuildResult,
    ScrapSettleRequest,
    ReportRow,
    StockBalanceListResponse,
    StockBalanceResponse,
    SupplementResponse,
    SupplementUpdate,
    TransferCreate,
    UnitResponse,
    UpstreamContractLookup,
    UserScopeResponse,
    UserScopeUpdate,
    VoidDocumentRequest,
    WarehouseCreate,
    WarehouseReportResponse,
    WarehouseResponse,
    WarehouseUpdate,
)
from app.services.warehouse.counts import WarehouseCountService
from app.services.warehouse.master import WarehouseMasterService
from app.services.warehouse.opening import WarehouseOpeningService
from app.services.warehouse.periods import WarehousePeriodService
from app.services.warehouse.posting import StockKey, WarehousePostingService, VOUCHER_FIELDS
from app.services.warehouse.queries import (
    WarehouseQueryService,
    serialize_balance,
    serialize_document,
    serialize_ledger,
)
from app.services.warehouse.scope import WarehouseScopeService
from app.services.warehouse.supplements import WarehouseSupplementService

router = APIRouter()


def _idempotency_key(header_key: str | None, body_key: str | None) -> str | None:
    return header_key or body_key


def _request_meta(request: Request) -> tuple[str | None, str | None]:
    return get_client_ip(request), request.headers.get("user-agent")


async def _assert_document_scope(
    db: AsyncSession, current_user: User, document
) -> None:
    warehouse_ids = {
        warehouse_id
        for line in document.lines
        for warehouse_id in (line.source_warehouse_id, line.target_warehouse_id)
        if warehouse_id
    }
    await WarehouseScopeService(db, current_user).assert_warehouses_access(
        warehouse_ids
    )


def _serialize_count(count) -> CountResponse:
    return CountResponse(
        id=count.id,
        count_no=count.count_no,
        warehouse_id=count.warehouse_id,
        location_id=count.location_id,
        project_id=count.project_id,
        counted_on=count.counted_on,
        snapshot_at=count.snapshot_at,
        snapshot_source=count.snapshot_source,
        status=count.status,
        description=count.description,
        void_reason=count.void_reason,
        review_notes=count.review_notes,
        created_by=count.created_by,
        confirmed_by=count.confirmed_by,
        reviewed_by=count.reviewed_by,
        voided_by=count.voided_by,
        created_at=count.created_at,
        confirmed_at=count.confirmed_at,
        reviewed_at=count.reviewed_at,
        voided_at=count.voided_at,
        reopened_from_id=count.reopened_from_id,
        adjustment_document_id=count.adjustment_document_id,
        warehouse_name=count.warehouse.name if count.warehouse else None,
        lines=[
            CountLineResponse(
                id=line.id,
                warehouse_id=line.warehouse_id,
                location_id=line.location_id,
                project_id=line.project_id,
                material_id=line.material_id,
                batch_no=line.batch_no or "",
                serial_no=line.serial_no or "",
                expiry_date=line.expiry_date,
                book_quantity=line.book_quantity,
                counted_quantity=line.counted_quantity,
                variance_reviewed=bool(line.variance_reviewed),
                variance_note=line.variance_note,
                material_code=line.material.code if line.material else None,
                material_name=line.material.name if line.material else None,
                material_unit=line.material.unit if line.material else None,
                location_name=line.location.name if line.location else None,
                project_name=line.project.name if line.project else None,
                adjustment_document_id=line.adjustment_document_id,
            )
            for line in count.lines
        ],
    )


@router.get("/warehouses", response_model=list[WarehouseResponse])
async def list_warehouses(
    include_inactive: bool = False,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    scope = WarehouseScopeService(db, current_user)
    allowed = await scope.authorized_warehouse_ids(for_posting=False)
    items = await WarehouseMasterService(db, current_user).list_warehouses(
        include_inactive=include_inactive
    )
    if allowed is not None:
        items = [item for item in items if item.id in allowed]
    return items


@router.post("/warehouses", response_model=WarehouseResponse, status_code=201)
async def create_warehouse(
    payload: WarehouseCreate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    warehouse = await WarehouseMasterService(db, current_user).create_warehouse(
        payload.model_dump()
    )
    await db.commit()
    return warehouse


@router.get("/warehouses/{warehouse_id}", response_model=WarehouseResponse)
async def get_warehouse(
    warehouse_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    await WarehouseScopeService(db, current_user).assert_warehouse_access(warehouse_id)
    return await WarehouseMasterService(db, current_user).get_warehouse(warehouse_id)


@router.put("/warehouses/{warehouse_id}", response_model=WarehouseResponse)
async def update_warehouse(
    warehouse_id: int,
    payload: WarehouseUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    warehouse = await WarehouseMasterService(db, current_user).update_warehouse(
        warehouse_id, payload.model_dump(exclude_unset=True)
    )
    await db.commit()
    return warehouse


@router.delete("/warehouses/{warehouse_id}", status_code=204)
async def delete_warehouse(
    warehouse_id: int,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    # 主数据生命周期：有业务的库房服务层拒绝物理删除（仅可停用），无业务方可删除
    await WarehouseMasterService(db, current_user).delete_warehouse(warehouse_id)
    await db.commit()


@router.get(
    "/warehouses/{warehouse_id}/locations", response_model=list[LocationResponse]
)
async def list_locations(
    warehouse_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    await WarehouseScopeService(db, current_user).assert_warehouse_access(warehouse_id)
    warehouse = await WarehouseMasterService(db, current_user).get_warehouse(
        warehouse_id
    )
    return warehouse.locations


@router.post(
    "/warehouses/{warehouse_id}/locations",
    response_model=LocationResponse,
    status_code=201,
)
async def create_location(
    warehouse_id: int,
    payload: LocationCreate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    location = await WarehouseMasterService(db, current_user).create_location(
        warehouse_id, payload.model_dump()
    )
    await db.commit()
    return location


@router.put("/locations/{location_id}", response_model=LocationResponse)
async def update_location(
    location_id: int,
    payload: LocationUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    location = await WarehouseMasterService(db, current_user).update_location(
        location_id, payload.model_dump(exclude_unset=True)
    )
    await db.commit()
    return location


@router.delete("/locations/{location_id}", status_code=204)
async def delete_location(
    location_id: int,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    # 有业务的货位服务层拒绝物理删除（仅可停用）
    await WarehouseMasterService(db, current_user).delete_location(location_id)
    await db.commit()


@router.get("/upstream-contracts", response_model=list[UpstreamContractLookup])
async def lookup_upstream_contracts(
    q: str | None = None,
    serial_number: int | None = None,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    contracts = await WarehouseMasterService(
        db, current_user
    ).lookup_upstream_contracts(q=q, serial_number=serial_number)
    return [
        UpstreamContractLookup(
            id=item.id,
            serial_number=item.serial_number,
            contract_name=item.contract_name,
            company_category=item.company_category,
            party_a_name=item.party_a_name,
        )
        for item in contracts
    ]


@router.get("/projects", response_model=list[ProjectResponse])
async def list_projects(
    q: str | None = None,
    include_inactive: bool = False,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseMasterService(db, current_user).list_projects(
        include_inactive=include_inactive, q=q
    )


@router.post("/projects", response_model=ProjectResponse, status_code=201)
async def create_project(
    payload: ProjectCreate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    project = await WarehouseMasterService(db, current_user).create_project(
        payload.model_dump()
    )
    await db.commit()
    return project


@router.get("/projects/{project_id}", response_model=ProjectResponse)
async def get_project(
    project_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseMasterService(db, current_user).get_project(project_id)


@router.put("/projects/{project_id}", response_model=ProjectResponse)
async def update_project(
    project_id: int,
    payload: ProjectUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    project = await WarehouseMasterService(db, current_user).update_project(
        project_id, payload.model_dump(exclude_unset=True)
    )
    await db.commit()
    return project


@router.get("/materials/search", response_model=list[MaterialResponse])
async def search_materials(
    q: str = Query(..., min_length=1),
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseMasterService(db, current_user).list_materials(q=q, limit=30)


@router.get("/materials", response_model=list[MaterialResponse])
async def list_materials(
    q: str | None = None,
    category: str | None = None,
    supply_type: str | None = None,
    condition: str | None = None,
    include_inactive: bool = False,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseMasterService(db, current_user).list_materials(
        q=q,
        category=category,
        supply_type=supply_type,
        condition=condition,
        include_inactive=include_inactive,
        limit=200,
    )


@router.post("/materials", response_model=MaterialResponse, status_code=201)
async def create_material(
    payload: MaterialCreate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MATERIALS)
    ),
    db: AsyncSession = Depends(get_db),
):
    material = await WarehouseMasterService(db, current_user).create_material(
        payload.model_dump()
    )
    await db.commit()
    return material


@router.get("/materials/{material_id}", response_model=MaterialResponse)
async def get_material(
    material_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseMasterService(db, current_user).get_material(material_id)


@router.put("/materials/{material_id}", response_model=MaterialResponse)
async def update_material(
    material_id: int,
    payload: MaterialUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MATERIALS)
    ),
    db: AsyncSession = Depends(get_db),
):
    material = await WarehouseMasterService(db, current_user).update_material(
        material_id, payload.model_dump(exclude_unset=True)
    )
    await db.commit()
    return material


@router.get("/materials/{material_id}/qr", response_model=QrPayloadResponse)
async def material_qr(
    material_id: int,
    request: Request,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    material = await WarehouseMasterService(db, current_user).get_material(material_id)
    origin = str(request.base_url).rstrip("/")
    payload = f"{origin}/m/warehouse/materials/{material.id}"
    return QrPayloadResponse(
        id=material.id,
        kind="material",
        payload=payload,
        label=f"{material.code} {material.name}",
    )


@router.get("/locations/{location_id}/qr", response_model=QrPayloadResponse)
async def location_qr(
    location_id: int,
    request: Request,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    location = await WarehouseMasterService(db, current_user).get_location(location_id)
    await WarehouseScopeService(db, current_user).assert_warehouse_access(
        location.warehouse_id
    )
    origin = str(request.base_url).rstrip("/")
    payload = f"{origin}/m/warehouse/locations/{location.id}"
    return QrPayloadResponse(
        id=location.id,
        kind="location",
        payload=payload,
        label=f"{location.code} {location.name}",
    )


@router.get("/locations/{location_id}", response_model=LocationResponse)
async def get_location(
    location_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    location = await WarehouseMasterService(db, current_user).get_location(location_id)
    await WarehouseScopeService(db, current_user).assert_warehouse_access(
        location.warehouse_id
    )
    return location


@router.get("/users/{user_id}/warehouse-scopes", response_model=list[UserScopeResponse])
async def list_user_scopes(
    user_id: int,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    scopes = await WarehouseScopeService(db, current_user).list_scopes(user_id)
    return [
        UserScopeResponse(
            user_id=scope.user_id,
            warehouse_id=scope.warehouse_id,
            warehouse_code=scope.warehouse.code if scope.warehouse else None,
            warehouse_name=scope.warehouse.name if scope.warehouse else None,
            is_default=scope.is_default,
        )
        for scope in scopes
    ]


@router.put("/users/{user_id}/warehouse-scopes", response_model=list[UserScopeResponse])
async def replace_user_scopes(
    user_id: int,
    payload: UserScopeUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_MASTER)
    ),
    db: AsyncSession = Depends(get_db),
):
    scopes = await WarehouseScopeService(db, current_user).replace_scopes(
        user_id, [item.model_dump() for item in payload.scopes]
    )
    await db.commit()
    return [
        UserScopeResponse(
            user_id=scope.user_id,
            warehouse_id=scope.warehouse_id,
            warehouse_code=scope.warehouse.code if scope.warehouse else None,
            warehouse_name=scope.warehouse.name if scope.warehouse else None,
            is_default=scope.is_default,
        )
        for scope in scopes
    ]


@router.get("/me/warehouse-scopes", response_model=list[UserScopeResponse])
async def my_warehouse_scopes(
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    scope = WarehouseScopeService(db, current_user)
    if scope.can_view_all():
        warehouses = await WarehouseMasterService(db, current_user).list_warehouses()
        default_id = await scope.default_warehouse_id()
        return [
            UserScopeResponse(
                user_id=current_user.id,
                warehouse_id=item.id,
                warehouse_code=item.code,
                warehouse_name=item.name,
                is_default=item.id == default_id,
            )
            for item in warehouses
        ]
    scopes = await scope.list_scopes(current_user.id)
    return [
        UserScopeResponse(
            user_id=item.user_id,
            warehouse_id=item.warehouse_id,
            warehouse_code=item.warehouse.code if item.warehouse else None,
            warehouse_name=item.warehouse.name if item.warehouse else None,
            is_default=item.is_default,
        )
        for item in scopes
        if item.warehouse and item.warehouse.is_active
    ]


@router.get("/stock-balances", response_model=StockBalanceListResponse)
async def list_stock_balances(
    warehouse_id: int | None = None,
    location_id: int | None = None,
    project_id: int | None = None,
    material_id: int | None = None,
    supply_type: str | None = None,
    condition: str | None = None,
    category: str | None = None,
    q: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    include_zero: bool = False,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    if warehouse_id:
        await WarehouseScopeService(db, current_user).assert_warehouse_access(
            warehouse_id
        )
    items, total = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_balances(
        warehouse_id=warehouse_id,
        location_id=location_id,
        project_id=project_id,
        material_id=material_id,
        supply_type=supply_type,
        condition=condition,
        category=category,
        q=q,
        page=page,
        page_size=page_size,
        include_zero=include_zero,
    )
    return StockBalanceListResponse(
        items=[serialize_balance(item) for item in items],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get(
    "/stock-balances/materials/{material_id}",
    response_model=list[StockBalanceResponse],
)
async def list_material_balances(
    material_id: int,
    warehouse_id: int | None = None,
    location_id: int | None = None,
    project_id: int | None = None,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    items, _ = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_balances(
        warehouse_id=warehouse_id,
        location_id=location_id,
        project_id=project_id,
        material_id=material_id,
        page=1,
        page_size=200,
        include_zero=True,
    )
    return [serialize_balance(item) for item in items]


@router.get("/stock-balances/available", response_model=StockBalanceResponse)
async def get_available_quantity(
    warehouse_id: int,
    location_id: int,
    project_id: int,
    material_id: int,
    batch_no: str | None = None,
    serial_no: str | None = None,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    await WarehouseScopeService(db, current_user).assert_warehouse_access(warehouse_id)
    service = WarehousePostingService(db, current_user)
    quantity = await service.get_available(
        StockKey(
            warehouse_id,
            location_id,
            project_id,
            material_id,
            batch_no or "",
            serial_no or "",
        )
    )
    return StockBalanceResponse(
        warehouse_id=warehouse_id,
        location_id=location_id,
        project_id=project_id,
        material_id=material_id,
        batch_no=batch_no,
        serial_no=serial_no,
        quantity=quantity,
        minimum_stock=0,
        version=0,
    )


@router.post("/stock-balances/rebuild", response_model=RebuildResult)
async def rebuild_balances(
    payload: RebuildRequest | None = None,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    body = payload or RebuildRequest()
    if body.repair and not has_permission(
        current_user, Permission.REPAIR_WAREHOUSE_BALANCES
    ):
        raise AppException(
            error_code=ErrorCode.INSUFFICIENT_PERMISSIONS,
            message="只有库房管理员可以执行库存对账修复",
            status_code=403,
        )
    result = await WarehousePostingService(db, current_user).rebuild_balances(
        repair=body.repair, reason=body.reason
    )
    await db.commit()
    return RebuildResult(**result)


@router.get("/ledger", response_model=LedgerListResponse)
async def list_ledger(
    warehouse_id: int | None = None,
    location_id: int | None = None,
    project_id: int | None = None,
    material_id: int | None = None,
    document_type: str | None = None,
    business_type: str | None = None,
    category: str | None = None,
    supply_type: str | None = None,
    condition: str | None = None,
    handler: str | None = None,
    q: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    items, total = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_ledger(
        warehouse_id=warehouse_id,
        location_id=location_id,
        project_id=project_id,
        material_id=material_id,
        document_type=document_type,
        business_type=business_type,
        category=category,
        supply_type=supply_type,
        condition=condition,
        handler=handler,
        q=q,
        start_date=start_date,
        end_date=end_date,
        page=page,
        page_size=page_size,
    )
    return LedgerListResponse(
        items=[serialize_ledger(item) for item in items],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/units", response_model=list[UnitResponse])
async def list_units(
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_units()


@router.get("/units/convert")
async def convert_unit(
    from_unit: str,
    to_unit: str,
    quantity: Decimal,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    converted = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).convert_quantity(from_code=from_unit, to_code=to_unit, quantity=quantity)
    return {"from_unit": from_unit, "to_unit": to_unit, "quantity": str(converted)}


@router.get("/reports/{kind}", response_model=WarehouseReportResponse)
async def get_report(
    kind: str,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    if kind not in {
        "consumption",
        "turnover",
        "low-stock",
        "idle",
        "expiring",
        "movement",
    }:
        raise ValidationError(message=f"未知报表类型: {kind}")
    items = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).report(kind)
    return WarehouseReportResponse(
        kind=kind, items=[ReportRow(**item) for item in items]
    )


@router.get("/documents", response_model=DocumentListResponse)
async def list_documents(
    document_type: str | None = None,
    status: str | None = None,
    business_type: str | None = None,
    warehouse_id: int | None = None,
    material_id: int | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    items, total = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_documents(
        document_type=document_type,
        status=status,
        business_type=business_type,
        warehouse_id=warehouse_id,
        material_id=material_id,
        start_date=start_date,
        end_date=end_date,
        page=page,
        page_size=page_size,
    )
    return DocumentListResponse(
        items=[serialize_document(item) for item in items],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/documents/{document_id}", response_model=DocumentResponse)
async def get_document(
    document_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    document = await WarehousePostingService(db, current_user).get_document(document_id)
    await _assert_document_scope(db, current_user, document)
    return serialize_document(document)


@router.post("/documents/{document_id}/void", response_model=DocumentResponse)
async def void_document(
    document_id: int,
    payload: VoidDocumentRequest,
    request: Request,
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    current_user: User = Depends(
        require_permission(Permission.VOID_WAREHOUSE_DOCUMENT)
    ),
    db: AsyncSession = Depends(get_db),
):
    ip_address, user_agent = _request_meta(request)
    try:
        document = await WarehousePostingService(db, current_user).void_document(
            document_id,
            payload.reason,
            idempotency_key=_idempotency_key(idempotency_key, payload.idempotency_key),
            ip_address=ip_address,
            user_agent=user_agent,
        )
    except IntegrityError as exc:
        await db.rollback()
        if "uq_warehouse_document_reversed_document" not in str(exc.orig):
            raise
        raise AppException(
            error_code=ErrorCode.WAREHOUSE_CONFLICT,
            message="单据已被其他请求冲销",
            status_code=409,
        ) from exc
    await db.commit()
    return serialize_document(document)


@router.get(
    "/documents/{document_id}/supplement", response_model=list[SupplementResponse]
)
async def get_supplements(
    document_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    document = await WarehousePostingService(db, current_user).get_document(document_id)
    await _assert_document_scope(db, current_user, document)
    return document.supplements


@router.put("/documents/{document_id}/supplement", response_model=SupplementResponse)
async def upsert_supplement(
    document_id: int,
    payload: SupplementUpdate,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_SUPPLEMENTS)
    ),
    db: AsyncSession = Depends(get_db),
):
    row = await WarehouseSupplementService(db, current_user).upsert(
        document_id, payload.model_dump()
    )
    await db.commit()
    await db.refresh(row)
    return row


@router.post("/inbounds", response_model=DocumentResponse, status_code=201)
async def create_inbound(
    payload: InboundCreate,
    request: Request,
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    current_user: User = Depends(require_permission(Permission.POST_WAREHOUSE_INBOUND)),
    db: AsyncSession = Depends(get_db),
):
    ip_address, user_agent = _request_meta(request)
    document = await WarehousePostingService(db, current_user).post_inbound(
        warehouse_id=payload.warehouse_id,
        location_id=payload.location_id,
        project_id=payload.project_id,
        occurred_on=payload.occurred_on,
        handler=payload.handler,
        business_type=payload.business_type.value,
        lines=[line.model_dump() for line in payload.lines],
        reference_no=payload.reference_no,
        description=payload.description,
        delivery_note_file=payload.delivery_note_file,
        delivery_note_file_name=payload.delivery_note_file_name,
        voucher=payload.model_dump(include=set(VOUCHER_FIELDS)),
        idempotency_key=_idempotency_key(idempotency_key, payload.idempotency_key),
        ip_address=ip_address,
        user_agent=user_agent,
    )
    await db.commit()
    return serialize_document(document)


@router.post("/outbounds", response_model=DocumentResponse, status_code=201)
async def create_outbound(
    payload: OutboundCreate,
    request: Request,
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    current_user: User = Depends(
        require_permission(Permission.POST_WAREHOUSE_OUTBOUND)
    ),
    db: AsyncSession = Depends(get_db),
):
    ip_address, user_agent = _request_meta(request)
    document = await WarehousePostingService(db, current_user).post_outbound(
        warehouse_id=payload.warehouse_id,
        location_id=payload.location_id,
        project_id=payload.project_id,
        occurred_on=payload.occurred_on,
        handler=payload.handler,
        business_type=payload.business_type.value,
        lines=[line.model_dump() for line in payload.lines],
        reference_no=payload.reference_no,
        description=payload.description,
        scrap_basis_file=payload.scrap_basis_file,
        scrap_basis_file_name=payload.scrap_basis_file_name,
        voucher=payload.model_dump(include=set(VOUCHER_FIELDS)),
        idempotency_key=_idempotency_key(idempotency_key, payload.idempotency_key),
        ip_address=ip_address,
        user_agent=user_agent,
    )
    await db.commit()
    return serialize_document(document)


@router.post("/transfers", response_model=DocumentResponse, status_code=201)
async def create_transfer(
    payload: TransferCreate,
    request: Request,
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    current_user: User = Depends(
        require_permission(Permission.POST_WAREHOUSE_TRANSFER)
    ),
    db: AsyncSession = Depends(get_db),
):
    ip_address, user_agent = _request_meta(request)
    document = await WarehousePostingService(db, current_user).post_transfer(
        occurred_on=payload.occurred_on,
        handler=payload.handler,
        reference_no=payload.reference_no,
        lines=[line.model_dump() for line in payload.lines],
        description=payload.description,
        idempotency_key=_idempotency_key(idempotency_key, payload.idempotency_key),
        ip_address=ip_address,
        user_agent=user_agent,
    )
    await db.commit()
    return serialize_document(document)


@router.post("/counts", response_model=CountResponse, status_code=201)
async def create_count(
    payload: CountCreate,
    current_user: User = Depends(require_permission(Permission.ENTER_WAREHOUSE_COUNT)),
    db: AsyncSession = Depends(get_db),
):
    service = WarehouseCountService(db, current_user)
    try:
        count = await service.create_count(
            warehouse_id=payload.warehouse_id,
            location_id=payload.location_id,
            project_id=payload.project_id,
            counted_on=payload.counted_on,
            description=payload.description,
            idempotency_key=payload.idempotency_key,
        )
    except IntegrityError as exc:
        await db.rollback()
        if "uq_warehouse_count_idempotency" not in str(exc.orig):
            raise
        count = await service.create_count(
            warehouse_id=payload.warehouse_id,
            location_id=payload.location_id,
            project_id=payload.project_id,
            counted_on=payload.counted_on,
            description=payload.description,
            idempotency_key=payload.idempotency_key,
        )
    await db.commit()
    return _serialize_count(count)


@router.get("/counts", response_model=CountListResponse)
async def list_counts(
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    items = await WarehouseCountService(db, current_user).list_counts()
    return CountListResponse(
        items=[
            CountResponse(
                id=item.id,
                count_no=item.count_no,
                warehouse_id=item.warehouse_id,
                location_id=item.location_id,
                project_id=item.project_id,
                counted_on=item.counted_on,
                snapshot_at=item.snapshot_at,
                snapshot_source=item.snapshot_source,
                status=item.status,
                description=item.description,
                void_reason=item.void_reason,
                review_notes=item.review_notes,
                created_by=item.created_by,
                confirmed_by=item.confirmed_by,
                reviewed_by=item.reviewed_by,
                voided_by=item.voided_by,
                created_at=item.created_at,
                confirmed_at=item.confirmed_at,
                reviewed_at=item.reviewed_at,
                voided_at=item.voided_at,
                reopened_from_id=item.reopened_from_id,
                adjustment_document_id=item.adjustment_document_id,
                warehouse_name=item.warehouse.name if item.warehouse else None,
                lines=[],
            )
            for item in items
        ],
        total=len(items),
        page=1,
        page_size=len(items) or 1,
    )


@router.get("/counts/{count_id}", response_model=CountResponse)
async def get_count(
    count_id: int,
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).get_count(count_id)
    await WarehouseScopeService(db, current_user).assert_warehouse_access(
        count.warehouse_id
    )
    return _serialize_count(count)


@router.put("/counts/{count_id}/lines", response_model=CountResponse)
async def update_count_lines(
    count_id: int,
    payload: CountLinesUpdate,
    current_user: User = Depends(require_permission(Permission.ENTER_WAREHOUSE_COUNT)),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).update_lines(
        count_id, [line.model_dump(exclude_none=True) for line in payload.lines]
    )
    await db.commit()
    return _serialize_count(count)


@router.post("/counts/{count_id}/confirm", response_model=CountResponse)
async def confirm_count(
    count_id: int,
    current_user: User = Depends(
        require_permission(Permission.CONFIRM_WAREHOUSE_COUNT)
    ),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).confirm(count_id)
    await db.commit()
    return _serialize_count(count)


@router.post("/counts/{count_id}/review", response_model=CountResponse)
async def review_count(
    count_id: int,
    payload: CountActionRequest | None = None,
    current_user: User = Depends(
        require_permission(Permission.CONFIRM_WAREHOUSE_COUNT)
    ),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).review(
        count_id, notes=(payload.notes if payload else None)
    )
    await db.commit()
    return _serialize_count(count)


@router.post("/counts/{count_id}/void", response_model=CountResponse)
async def void_count(
    count_id: int,
    payload: CountActionRequest,
    current_user: User = Depends(
        require_permission(Permission.CONFIRM_WAREHOUSE_COUNT)
    ),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).void(
        count_id, payload.reason or ""
    )
    await db.commit()
    return _serialize_count(count)


@router.post("/counts/{count_id}/reopen", response_model=CountResponse, status_code=201)
async def reopen_count(
    count_id: int,
    payload: CountActionRequest | None = None,
    current_user: User = Depends(
        require_permission(Permission.CONFIRM_WAREHOUSE_COUNT)
    ),
    db: AsyncSession = Depends(get_db),
):
    count = await WarehouseCountService(db, current_user).reopen(
        count_id, description=(payload.description if payload else None)
    )
    await db.commit()
    return _serialize_count(count)


@router.get("/periods", response_model=list[PeriodResponse])
async def list_periods(
    current_user: User = Depends(
        require_permission(Permission.VIEW_WAREHOUSE_INVENTORY)
    ),
    db: AsyncSession = Depends(get_db),
):
    return await WarehousePeriodService(db, current_user).list_periods()


@router.post("/periods/close", response_model=PeriodResponse)
async def close_period(
    payload: PeriodActionRequest,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_PERIODS)
    ),
    db: AsyncSession = Depends(get_db),
):
    period = await WarehousePeriodService(db, current_user).close_period(
        payload.year, payload.month
    )
    await db.commit()
    return period


@router.post("/periods/reopen", response_model=PeriodResponse)
async def reopen_period(
    payload: PeriodActionRequest,
    current_user: User = Depends(
        require_permission(Permission.MANAGE_WAREHOUSE_PERIODS)
    ),
    db: AsyncSession = Depends(get_db),
):
    period = await WarehousePeriodService(db, current_user).reopen_period(
        payload.year, payload.month, payload.reason or ""
    )
    await db.commit()
    return period


@router.post("/documents/{document_id}/scrap-settle", response_model=DocumentResponse)
async def settle_scrap(
    document_id: int,
    payload: ScrapSettleRequest,
    current_user: User = Depends(
        require_permission(Permission.SETTLE_WAREHOUSE_SCRAP)
    ),
    db: AsyncSession = Depends(get_db),
):
    document = await WarehousePostingService(db, current_user).settle_scrap(
        document_id,
        **payload.model_dump(),
    )
    await db.commit()
    return serialize_document(document)


@router.get("/opening-entries/template.xlsx")
async def download_opening_template(
    current_user: User = Depends(require_permission(Permission.IMPORT_WAREHOUSE_DATA)),
    db: AsyncSession = Depends(get_db),
):
    workbook = await WarehouseOpeningService(db, current_user).build_template()
    return _excel_response(workbook, "期初材料录入表.xlsx")


@router.post("/opening-entries/import", response_model=OpeningImportResult)
async def import_opening_entries(
    request: Request,
    file: UploadFile = File(...),
    current_user: User = Depends(require_permission(Permission.IMPORT_WAREHOUSE_DATA)),
    db: AsyncSession = Depends(get_db),
):
    ip_address, user_agent = _request_meta(request)
    content = await file.read()
    result = await WarehouseOpeningService(db, current_user).import_workbook(
        content,
        filename=file.filename,
        ip_address=ip_address,
        user_agent=user_agent,
    )
    await db.commit()
    return result


def _excel_response(workbook: Workbook, filename: str) -> StreamingResponse:
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
        },
    )


@router.get("/exports/materials.xlsx")
async def export_materials(
    current_user: User = Depends(require_permission(Permission.EXPORT_WAREHOUSE_DATA)),
    db: AsyncSession = Depends(get_db),
):
    items = await WarehouseMasterService(db, current_user).list_materials(
        include_inactive=True, limit=10000
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "物资档案"
    sheet.append(
        [
            "编码",
            "旧编码",
            "名称",
            "品牌",
            "规格",
            "单位",
            "类别",
            "供应",
            "成色",
            "状态",
        ]
    )
    for item in items:
        sheet.append(
            [
                item.code,
                item.legacy_code,
                item.name,
                item.brand,
                item.specification,
                item.unit,
                item.category,
                item.supply_type,
                item.condition,
                "启用" if item.is_active else "归档",
            ]
        )
    return _excel_response(workbook, "warehouse-materials.xlsx")


@router.get("/exports/stock.xlsx")
async def export_stock(
    current_user: User = Depends(require_permission(Permission.EXPORT_WAREHOUSE_DATA)),
    db: AsyncSession = Depends(get_db),
):
    items, _ = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_balances(page=1, page_size=10000, include_zero=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存余额"
    sheet.append(
        ["库房", "货位", "项目", "物资编码", "物资名称", "数量", "单位", "供应", "成色"]
    )
    for item in items:
        sheet.append(
            [
                item.warehouse.name if item.warehouse else item.warehouse_id,
                item.location.name if item.location else item.location_id,
                item.project.name if item.project else item.project_id,
                item.material.code if item.material else item.material_id,
                item.material.name if item.material else "",
                float(item.quantity),
                item.material.unit if item.material else "",
                item.material.supply_type if item.material else "",
                item.material.condition if item.material else "",
            ]
        )
    return _excel_response(workbook, "warehouse-stock.xlsx")


@router.get("/exports/ledger.xlsx")
async def export_ledger(
    current_user: User = Depends(require_permission(Permission.EXPORT_WAREHOUSE_DATA)),
    db: AsyncSession = Depends(get_db),
):
    items, _ = await WarehouseQueryService(
        db, WarehouseScopeService(db, current_user)
    ).list_ledger(page=1, page_size=10000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存流水"
    sheet.append(
        [
            "日期",
            "单号",
            "类型",
            "库房",
            "货位",
            "项目",
            "物资编码",
            "物资名称",
            "数量变化",
        ]
    )
    for item in items:
        sheet.append(
            [
                item.occurred_on.isoformat(),
                item.document.document_no if item.document else item.document_id,
                item.document.document_type if item.document else "",
                item.warehouse.name if item.warehouse else item.warehouse_id,
                item.location.name if item.location else item.location_id,
                item.project.name if item.project else item.project_id,
                item.material.code if item.material else item.material_id,
                item.material.name if item.material else "",
                float(item.quantity_delta),
            ]
        )
    return _excel_response(workbook, "warehouse-ledger.xlsx")
